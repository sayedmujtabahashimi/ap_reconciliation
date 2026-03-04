# -*- coding: utf-8 -*-
from odoo import models, fields, api
from odoo.exceptions import UserError
import base64
import io
import logging

_logger = logging.getLogger(__name__)

try:
    import openpyxl
except ImportError:
    openpyxl = None


class ApManualReconciliation(models.Model):
    """
    Manual Reconciliation — compare two Excel files directly.
    Both files must follow the same template:
    Number | Customer | Total | Invoice Date
    """
    _name        = 'ap.manual.reconciliation'
    _description = 'AP Manual Reconciliation'
    _order       = 'id desc'
    _rec_name    = 'name'

    name = fields.Char(string='Reference', required=True, default='New Manual Reconciliation')

    # File A — first Excel
    file_a        = fields.Binary(string='Afghan Post File (Excel)', attachment=True)
    file_a_name   = fields.Char(string='Afghan Post File Name')
    file_a_label  = fields.Char(string='File A Label', default='Afghan Post',
                                 help='Label shown in columns e.g. "Odoo Export"')

    # File B — second Excel
    file_b        = fields.Binary(string='HesabPay File (Excel)', attachment=True)
    file_b_name   = fields.Char(string='HesabPay File Name')
    file_b_label  = fields.Char(string='File B Label', default='HesabPay',
                                 help='Label shown in columns e.g. "HesabPay"')

    state = fields.Selection([
        ('draft',      'Draft'),
        ('reconciled', 'Reconciled'),
    ], default='draft', required=True)

    # Lines
    line_ids = fields.One2many('ap.manual.reconciliation.line', 'reconciliation_id', string='All Lines')
    mismatch_line_ids = fields.One2many(
        'ap.manual.reconciliation.line', 'reconciliation_id',
        string='Mismatches',
        domain=[
            '|', ('name_mismatch', '=', True),
            '|', ('amount_mismatch', '=', True),
                 ('found_in_b', '=', False),
        ],
    )

    # Summary
    total_lines     = fields.Integer(compute='_compute_summary', store=True)
    matched_lines   = fields.Integer(compute='_compute_summary', store=True)
    mismatch_lines  = fields.Integer(compute='_compute_summary', store=True)
    not_found_lines = fields.Integer(compute='_compute_summary', store=True)

    @api.depends('line_ids.found_in_b', 'line_ids.name_mismatch', 'line_ids.amount_mismatch')
    def _compute_summary(self):
        for rec in self:
            lines = rec.line_ids
            rec.total_lines     = len(lines)
            rec.not_found_lines = len(lines.filtered(lambda l: not l.found_in_b))
            rec.mismatch_lines  = len(lines.filtered(
                lambda l: l.found_in_b and (l.name_mismatch or l.amount_mismatch)
            ))
            rec.matched_lines = len(lines.filtered(
                lambda l: l.found_in_b and not l.name_mismatch and not l.amount_mismatch
            ))

    # ── Actions ───────────────────────────────────────────────────────────────

    def action_compare(self):
        self.ensure_one()
        if not self.file_a:
            raise UserError('Please upload File A.')
        if not self.file_b:
            raise UserError('Please upload File B.')
        if not openpyxl:
            raise UserError('openpyxl is not installed. Run: pip3 install openpyxl')

        data_a = self._parse_excel(self.file_a)
        data_b = self._parse_excel(self.file_b)

        self.line_ids.unlink()

        vals_list = []
        for number, row_a in data_a.items():
            row_b = data_b.get(number)
            if row_b:
                diff = row_a['total'] - row_b['total']
                vals_list.append({
                    'reconciliation_id': self.id,
                    'invoice_number':    number,
                    'a_customer':        row_a['customer'],
                    'a_date':            row_a['date'],
                    'a_total':           row_a['total'],
                    'b_customer':        row_b['customer'],
                    'b_date':            row_b['date'],
                    'b_total':           row_b['total'],
                    'found_in_b':        True,
                    'amount_difference': diff,
                    'name_mismatch':     row_a['customer'].strip().lower() != row_b['customer'].strip().lower(),
                    'amount_mismatch':   abs(diff) > 0.01,
                })
            else:
                vals_list.append({
                    'reconciliation_id': self.id,
                    'invoice_number':    number,
                    'a_customer':        row_a['customer'],
                    'a_date':            row_a['date'],
                    'a_total':           row_a['total'],
                    'b_customer':        '',
                    'b_date':            False,
                    'b_total':           0.0,
                    'found_in_b':        False,
                    'amount_difference': 0.0,
                    'name_mismatch':     False,
                    'amount_mismatch':   False,
                })

        if vals_list:
            self.env['ap.manual.reconciliation.line'].create(vals_list)

        self.write({'state': 'reconciled'})
        return self._reload()

    def action_export_report(self):
        self.ensure_one()
        if not openpyxl:
            raise UserError('openpyxl is not installed.')

        from openpyxl.styles import PatternFill, Font, Alignment
        wb  = openpyxl.Workbook()
        ws  = wb.active
        ws.title = 'Manual Reconciliation'

        hdr_fill      = PatternFill('solid', fgColor='4472C4')
        mismatch_fill = PatternFill('solid', fgColor='FFCCCC')
        ok_fill       = PatternFill('solid', fgColor='CCFFCC')
        nf_fill       = PatternFill('solid', fgColor='FFF2CC')
        hdr_font      = Font(bold=True, color='FFFFFF')

        label_a = self.file_a_label or 'File A'
        label_b = self.file_b_label or 'File B'

        headers = [
            'Invoice #',
            f'{label_a} Customer', f'{label_a} Date', f'{label_a} Total',
            f'{label_b} Customer', f'{label_b} Date', f'{label_b} Total',
            'Difference', 'Name Match?', 'Amount Match?', 'Result',
        ]
        ws.append(headers)
        for i, _ in enumerate(headers, 1):
            c = ws.cell(row=1, column=i)
            c.fill = hdr_fill
            c.font = hdr_font
            c.alignment = Alignment(horizontal='center')

        for line in self.line_ids:
            if not line.found_in_b:
                result, fill = 'NOT FOUND', nf_fill
            elif line.name_mismatch or line.amount_mismatch:
                result, fill = 'MISMATCH', mismatch_fill
            else:
                result, fill = 'OK', ok_fill

            ws.append([
                line.invoice_number,
                line.a_customer, str(line.a_date or ''), line.a_total,
                line.b_customer, str(line.b_date or ''), line.b_total,
                line.amount_difference,
                'NO' if line.name_mismatch  else 'YES',
                'NO' if line.amount_mismatch else 'YES',
                result,
            ])
            r = ws.max_row
            for col in range(1, len(headers) + 1):
                ws.cell(row=r, column=col).fill = fill

        for col in ws.columns:
            w = max((len(str(c.value or '')) for c in col), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(w + 4, 45)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        attach = self.env['ir.attachment'].create({
            'name':     f'{self.name}_report.xlsx',
            'type':     'binary',
            'datas':    base64.b64encode(buf.read()),
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        })
        return {
            'type':   'ir.actions.act_url',
            'url':    f'/web/content/{attach.id}?download=true',
            'target': 'self',
        }

    def action_reset(self):
        self.ensure_one()
        self.line_ids.unlink()
        self.write({
            'state':       'draft',
            'file_a':      False,
            'file_a_name': False,
            'file_b':      False,
            'file_b_name': False,
        })
        return self._reload()

    def _reload(self):
        return {
            'type':      'ir.actions.act_window',
            'res_model': 'ap.manual.reconciliation',
            'res_id':    self.id,
            'view_mode': 'form',
            'target':    'current',
        }

    def _parse_excel(self, file_field):
        """Parse an Excel file with columns: Number | Customer | Total | Invoice Date"""
        file_data = base64.b64decode(file_field)
        wb = openpyxl.load_workbook(io.BytesIO(file_data), data_only=True)
        ws = wb.active

        data = {}
        header_skipped = False
        for row in ws.iter_rows(values_only=True):
            if not header_skipped:
                header_skipped = True
                continue
            if not any(row):
                continue
            try:
                number   = str(row[0]).strip() if row[0] is not None else ''
                customer = str(row[1]).strip() if row[1] is not None else ''
                total    = float(row[2])        if row[2] is not None else 0.0
                date     = row[3]
                if hasattr(date, 'date'):
                    date = date.date()
                elif isinstance(date, str):
                    from datetime import datetime
                    for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%m/%d/%Y'):
                        try:
                            date = datetime.strptime(date, fmt).date()
                            break
                        except ValueError:
                            date = None
                if number:
                    data[number] = {'customer': customer, 'total': total, 'date': date}
            except Exception as e:
                _logger.warning('Skipping row: %s', e)
        return data


class ApManualReconciliationLine(models.Model):
    _name        = 'ap.manual.reconciliation.line'
    _description = 'AP Manual Reconciliation Line'
    _order       = 'invoice_number asc'

    reconciliation_id = fields.Many2one('ap.manual.reconciliation', ondelete='cascade', required=True)

    invoice_number = fields.Char(string='Invoice #')

    # File A side
    a_customer = fields.Char(string='AP Customer')
    a_date     = fields.Date(string='AP Date')
    a_total    = fields.Float(string='AP Total', digits=(16, 2))

    # File B side
    found_in_b = fields.Boolean(string='Found in HesabPay', default=False)
    b_customer = fields.Char(string='HP Customer')
    b_date     = fields.Date(string='HP Date')
    b_total    = fields.Float(string='HP Total', digits=(16, 2))

    # Difference
    amount_difference = fields.Float(string='Difference', digits=(16, 2))
    name_mismatch     = fields.Boolean(string='Name Mismatch',   default=False)
    amount_mismatch   = fields.Boolean(string='Amount Mismatch', default=False)
